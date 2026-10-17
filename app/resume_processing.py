"""Bounded resume format detection, extraction, and business limits."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol, cast

import filetype
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.config import Settings
from app.errors import AppError

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_SUFFIXES = {".pdf": PDF, ".docx": DOCX}
_WHITESPACE = re.compile(r"[ \t\f\v]+")
_BLANK_LINES = re.compile(r"\n{3,}")


@dataclass(frozen=True, slots=True)
class ResumeDocument:
    """Validated document text and the page model used for the P0 limit."""

    detected_mime: str
    page_count: int
    text: str
    page_map: list[dict[str, int]]


class FileTypeMatch(Protocol):
    """Minimal typed view of a ``filetype`` match."""

    mime: str


def sanitize_filename(filename: str | None) -> str:
    """Keep a display-only basename and reject paths or control characters."""

    candidate = Path(filename or "").name.strip()
    candidate = "".join(char for char in candidate if char.isprintable())
    if not candidate:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="请选择 PDF 或 DOCX 简历文件",
        )
    return candidate[:255]


def validate_file_identity(
    *,
    filename: str,
    declared_mime: str | None,
    content: bytes,
    settings: Settings,
) -> str:
    """Validate extension, declared MIME, size, and magic before persistence."""

    if not content:
        raise AppError(status_code=422, code="invalid_request", message="文件不能为空")
    if len(content) > settings.resume_max_bytes:
        raise AppError(status_code=413, code="file_too_large", message="文件不能超过 10 MB")

    suffix = Path(filename).suffix.lower()
    expected_mime = SUPPORTED_SUFFIXES.get(suffix)
    if expected_mime is None:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="仅支持 PDF 或 DOCX 简历",
        )
    if declared_mime and declared_mime.split(";", 1)[0].strip().lower() != expected_mime:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="文件类型与扩展名不一致",
        )
    detected = cast(FileTypeMatch | None, filetype.guess(content))
    if detected is None or detected.mime != expected_mime:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="文件内容不是有效的 PDF 或 DOCX",
        )
    return detected.mime


def extract_resume(
    content: bytes,
    *,
    detected_mime: str,
    settings: Settings,
) -> ResumeDocument:
    """Extract text with an explicit page model and stable user-facing failures."""

    if detected_mime == PDF:
        return _extract_pdf(content, settings)
    if detected_mime == DOCX:
        return _extract_docx(content, settings)
    raise AppError(
        status_code=422,
        code="unsupported_source",
        message="仅支持 PDF 或 DOCX 简历",
    )


def _normalize_text(value: str) -> str:
    lines = [_WHITESPACE.sub(" ", line).strip() for line in value.splitlines()]
    return _BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


def _extract_pdf(content: bytes, settings: Settings) -> ResumeDocument:
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if len(reader.pages) > settings.resume_max_pages:
            raise AppError(
                status_code=422,
                code="page_limit_exceeded",
                message="简历不能超过 10 页",
            )
        page_texts: list[str] = []
        page_map: list[dict[str, int]] = []
        cursor = 0
        for page in reader.pages:
            text = _normalize_text(page.extract_text() or "")
            if text:
                if page_texts:
                    cursor += 2
                start = cursor
                page_texts.append(text)
                cursor += len(text)
                page_map.append({"page": len(page_map) + 1, "start": start, "end": cursor})
    except AppError:
        raise
    except (PdfReadError, ValueError, TypeError, KeyError) as exc:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="无法读取该 PDF，请确认文件未损坏",
        ) from exc
    combined = "\n\n".join(page_texts)
    if not combined:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="扫描件或图片简历暂不支持",
        )
    return ResumeDocument(
        detected_mime=PDF,
        page_count=len(reader.pages),
        text=combined,
        page_map=page_map,
    )


def _extract_docx(content: bytes, settings: Settings) -> ResumeDocument:
    try:
        document = Document(io.BytesIO(content))
        text = _normalize_text("\n".join(paragraph.text for paragraph in document.paragraphs))
    except (ValueError, KeyError, TypeError) as exc:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="无法读取该 DOCX，请确认文件未损坏",
        ) from exc
    if not text:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="文档没有可提取的文本",
        )

    # DOCX does not expose settled printer pagination. Use deterministic
    # character capacity only as the P0 limit, not as a rendering claim.
    capacity = settings.resume_max_docx_characters
    page_count = max(1, (len(text) + capacity - 1) // capacity)
    if page_count > settings.resume_max_pages:
        raise AppError(
            status_code=422,
            code="page_limit_exceeded",
            message="简历内容超过 10 页上限",
        )
    page_map = [
        {
            "page": index + 1,
            "start": index * capacity,
            "end": min((index + 1) * capacity, len(text)),
        }
        for index in range(page_count)
    ]
    return ResumeDocument(
        detected_mime=DOCX,
        page_count=page_count,
        text=text,
        page_map=page_map,
    )
