"""Bounded PDF, DOCX, text, and Markdown extraction."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import cast

import filetype
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.config import Settings
from app.errors import AppError
from app.rag.contracts import ExtractedDocument

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT = "text/plain"
MARKDOWN = "text/markdown"
SUPPORTED_SUFFIXES = {
    ".pdf": PDF,
    ".docx": DOCX,
    ".txt": TEXT,
    ".md": MARKDOWN,
}
_WHITESPACE = re.compile(r"[ \t\f\v]+")
_BLANK_LINES = re.compile(r"\n{3,}")


def _sanitize_filename(filename: str | None) -> str:
    candidate = Path(filename or "").name.strip()
    candidate = "".join(char for char in candidate if char.isprintable())
    if not candidate:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="请选择支持的资料文件",
        )
    return candidate[:255]


def _normalize_text(value: str) -> str:
    lines = [_WHITESPACE.sub(" ", line).strip() for line in value.splitlines()]
    return _BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


def _validate_mime(
    *,
    filename: str,
    declared_mime: str | None,
    content: bytes,
) -> str:
    if not content:
        raise AppError(status_code=422, code="invalid_request", message="文件不能为空")
    suffix = Path(filename).suffix.lower()
    expected = SUPPORTED_SUFFIXES.get(suffix)
    if expected is None:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="仅支持 PDF、DOCX、TXT 或 Markdown 资料",
        )
    normalized_declared = (declared_mime or "").split(";", 1)[0].strip().lower()
    if expected in {TEXT, MARKDOWN}:
        if normalized_declared and normalized_declared not in {TEXT, MARKDOWN}:
            raise AppError(
                status_code=422,
                code="unsupported_source",
                message="文件类型与扩展名不一致",
            )
        return expected
    detected = cast(object | None, filetype.guess(content))
    detected_mime = getattr(detected, "mime", None)
    if detected_mime != expected:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="文件内容不是有效的 PDF 或 DOCX",
        )
    if normalized_declared and normalized_declared != expected:
        raise AppError(
            status_code=422,
            code="unsupported_source",
            message="文件类型与扩展名不一致",
        )
    return expected


def _bounded_text(text: str, settings: Settings) -> str:
    normalized = _normalize_text(text)
    if not normalized:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="文档没有可提取的文本",
        )
    if len(normalized) > settings.document_max_characters:
        raise AppError(
            status_code=422,
            code="document_too_large",
            message="文档正文超过长度上限",
        )
    return normalized


def _extract_pdf(content: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, TypeError, KeyError) as exc:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="无法读取该 PDF，请确认文件未损坏",
        ) from exc
    return "\n\n".join(pages), len(reader.pages)


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except (ValueError, KeyError, TypeError) as exc:
        raise AppError(
            status_code=422,
            code="text_extraction_unsupported",
            message="无法读取该 DOCX，请确认文件未损坏",
        ) from exc


def extract_document(
    content: bytes,
    *,
    filename: str,
    declared_mime: str | None,
    settings: Settings,
) -> ExtractedDocument:
    """Validate and extract one bounded private document."""

    safe_filename = _sanitize_filename(filename)
    if len(content) > settings.document_max_bytes:
        raise AppError(status_code=413, code="file_too_large", message="文件不能超过 10 MB")
    detected_mime = _validate_mime(
        filename=safe_filename,
        declared_mime=declared_mime,
        content=content,
    )
    page_count: int | None = None
    if detected_mime == PDF:
        text, page_count = _extract_pdf(content)
    elif detected_mime == DOCX:
        text = _extract_docx(content)
    else:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise AppError(
                status_code=422,
                code="text_extraction_unsupported",
                message="文本文件必须使用 UTF-8 编码",
            ) from exc
    return ExtractedDocument(
        detected_mime=detected_mime,
        text=_bounded_text(text, settings),
        page_count=page_count,
    )
